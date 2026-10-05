"""Add analytics using only the already authorized, filtered export records."""
from collections import defaultdict
from decimal import Decimal
from math import ceil

from django.utils import timezone
from django.utils.translation import get_language
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.axis import ChartLines
from openpyxl.chart.data_source import AxDataSource, StrRef
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import CharacterProperties, Paragraph, ParagraphProperties, RichTextProperties
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def add_stock_report_analytics(workbook, requests, filters):
    french = (get_language() or "en").startswith("fr")

    def tr(en, fr):
        return fr if french else en

    material_lines = 0
    materials = {}
    departments = {}
    months = defaultdict(lambda: [0, set()])
    missing_dates = set()
    for req in requests:
        dept = departments.setdefault(req.request_for_department_id,
                                      [req.request_for_department.name, set(), 0])
        dept[1].add(req.pk)
        month = req.finalized_at.strftime("%Y-%m") if req.finalized_at else None
        if month:
            months[month][1].add(req.pk)
        else:
            missing_dates.add(req.pk)
        for item in req.material_items.all():
            material_lines += 1
            dept[2] += 1
            material = materials.setdefault(item.material_id,
                [item.material.name, item.material.code, item.material.unit, Decimal("0")])
            material[3] += item.quantity
            if month:
                months[month][0] += 1

    top = sorted(materials.values(), key=lambda m: (-m[3], m[0], m[1]))[:10]
    depts = sorted(departments.values(), key=lambda d: (-len(d[1]), d[0]))
    dashboard = workbook.create_sheet(tr("Dashboard", "Tableau de bord"))
    graphs = workbook.create_sheet(tr("Graphs", "Graphiques"))
    navy = "17365D"
    for sheet, title in ((dashboard, tr("STOCK REPORT DASHBOARD", "TABLEAU DE BORD DU STOCK")),
                         (graphs, tr("STOCK REPORT — VISUAL ANALYSIS", "RAPPORT DE MATÉRIEL — ANALYSE VISUELLE"))):
        sheet.sheet_view.showGridLines = False
        last_col = 21 if sheet is graphs else 12
        for col in range(1, last_col + 1):
            sheet.column_dimensions[get_column_letter(col)].width = 10 if sheet is graphs else 14
        if sheet is graphs:
            sheet.column_dimensions["K"].width = 3
            for row in range(8, 45):
                sheet.row_dimensions[row].height = 18
        for row, text in ((1, "MICROCOM"), (2, title)):
            sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
            cell = sheet.cell(row, 1, text)
            cell.font = Font(size=20 if row == 1 else 16, bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=navy)
            sheet.row_dimensions[row].height = 30
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True

    start, end = (filters.get(k, "").strip() for k in ("date_from", "date_to"))
    period = f"{start or tr('No lower bound', 'Sans borne inférieure')} — {end or tr('No upper bound', 'Sans borne supérieure')}"
    if not start and not end:
        period = tr("All dates (no date filter)", "Toutes les dates (aucun filtre de date)")
    dept_filter = filters.get("department", "").strip()
    dept_name = departments.get(int(dept_filter), [None])[0] if dept_filter.isdigit() else None
    metadata = [
        tr("Report period (Date Needed): ", "Période (Date requise) : ") + period,
        tr("Department: ", "Département : ") + (dept_name or (f"ID {dept_filter}" if dept_filter else tr("All", "Tous"))),
        tr("Search: ", "Recherche : ") + (filters.get("q", "").strip() or tr("None", "Aucune")),
        tr("Generated At (UTC): ", "Généré le (UTC) : ") + timezone.now().strftime("%Y-%m-%d %H:%M"),
    ]
    for row, text in enumerate(metadata, 4):
        dashboard.merge_cells(start_row=row, start_column=1, end_row=row, end_column=12)
        dashboard.cell(row, 1, text).alignment = Alignment(wrap_text=True)
        dashboard.row_dimensions[row].height = 24
    for coordinate, end_col, text in (("A3", 10, metadata[0]), ("L3", 21, metadata[3]),
                                      ("A4", 10, metadata[1]), ("L4", 21, metadata[2])):
        cell = graphs[coordinate]
        graphs.merge_cells(start_row=cell.row, start_column=cell.column, end_row=cell.row, end_column=end_col)
        cell.value = text
        cell.font = Font(size=11, color=navy)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        graphs.row_dimensions[cell.row].height = 24
    labels = [tr("Total Approved Requests", "Total des demandes approuvées"),
              tr("Material Lines Issued", "Lignes d'articles sorties"),
              tr("Unique Materials", "Matériels distincts"),
              tr("Departments Served", "Départements servis")]
    for index, (label, value) in enumerate(zip(labels, (len(requests), material_lines, len(materials), len(departments)))):
        col = index * 3 + 1
        for row, content in ((9, label), (10, value)):
            dashboard.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + 2)
            cell = dashboard.cell(row, col, content)
            cell.fill = PatternFill("solid", fgColor="D9EAF7")
            cell.font = Font(size=12 if row == 9 else 24, bold=True, color=navy)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        dashboard.cell(10, col).number_format = "#,##0"
    dashboard.row_dimensions[9].height = 36
    dashboard.row_dimensions[10].height = 40
    note = tr("Material Report is the detailed source. Both analytical sheets use the same filtered records.\nTime charts use Approved Date. Quantities are kept per material/unit; no global quantity total.",
              "Le Rapport de matériel est la source détaillée. Les deux feuilles d’analyse utilisent les mêmes données filtrées.\nLes tendances utilisent la date d’approbation. Quantités par matériel/unité, sans total global.")
    if missing_dates:
        note += "\n" + tr("Requests without Approved Date excluded from time charts: ", "Demandes sans date d’approbation exclues des graphiques temporels : ") + str(len(missing_dates))
    for sheet, region, coordinate in ((dashboard, "A12:L13", "A12"), (graphs, "A5:U6", "A5")):
        sheet.merge_cells(region)
        cell = sheet[coordinate]
        cell.value = note
        cell.font = Font(size=11, color=navy)
        cell.fill = PatternFill("solid", fgColor="EDF3F8")
        cell.alignment = Alignment(wrap_text=True, vertical="center", indent=1)
        sheet.row_dimensions[cell.row].height = 30 if missing_dates else 24
        sheet.row_dimensions[cell.row + 1].height = 24
    empty = tr("No stock data available for the selected filters.", "Aucune donnée de stock disponible pour les filtres sélectionnés.")
    if not materials:
        for sheet in (dashboard, graphs):
            sheet.merge_cells("A15:L16")
            sheet["A15"] = empty
            sheet.print_area = "A1:U16" if sheet is graphs else "A1:L16"
        return

    def table(sheet, row, col, headers, data):
        for offset, header in enumerate(headers):
            cell = sheet.cell(row, col + offset, header)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=navy)
            cell.alignment = Alignment(wrap_text=True)
        if sheet is dashboard:
            sheet.row_dimensions[row].height = 32
        for r, values in enumerate(data, row + 1):
            for c, value in enumerate(values, col):
                cell = sheet.cell(r, c, value)
                # Treat names/codes as text even when they begin with '='.
                if isinstance(value, str):
                    cell.data_type = "s"
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                if isinstance(value, Decimal):
                    cell.number_format = "#,##0.##"

    material_label = tr("Material", "Matériel")
    quantity_label = tr("Quantity Issued", "Quantité sortie")
    dept_label = tr("Department", "Département")
    requests_label = tr("Approved Requests", "Demandes approuvées")
    lines_label = tr("Material Lines", "Lignes d'articles")
    dashboard["A15"] = tr("TOP MATERIALS", "PRINCIPAUX MATÉRIELS")
    table(dashboard, 16, 1, [material_label, tr("Code", "Code"), tr("Unit", "Unité"), quantity_label], top)
    dashboard["F15"] = tr("DEPARTMENT SUMMARY", "SYNTHÈSE PAR DÉPARTEMENT")
    table(dashboard, 16, 6, [dept_label, requests_label, lines_label],
          [(name, len(ids), lines) for name, ids, lines in depts])
    # Equal-width card groups, with room for both summary tables.
    for col, width in {"A": 28, "B": 16, "C": 12, "D": 18, "E": 12, "F": 26,
                       "G": 20, "H": 18, "I": 18, "J": 18, "K": 18, "L": 20}.items():
        dashboard.column_dimensions[col].width = width
    dashboard.print_area = f"A1:L{max(27, 16 + len(depts))}"

    def material_category(name, code, unit):
        label = name or code
        # Codes disambiguate duplicate display names without lengthening every label.
        if sum(m[0] == name for m in top) > 1 and code:
            label = f"{code} — {label}"
        return f"{label} ({unit})" if unit else label

    table(graphs, 1, 25, [material_label, quantity_label],
          [(material_category(name, code, unit), qty) for name, code, unit, qty in top])
    table(graphs, 1, 28, [dept_label, requests_label], [(d[0], len(d[1])) for d in depts])
    # The installed French catalog omits Django's month abbreviations. Keep
    # these labels explicitly bilingual, like the other analytical sheet text.
    month_names = tr(("Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"),
                     ("janv.", "févr.", "mars", "avr.", "mai", "juin", "juil.", "août", "sept.", "oct.", "nov.", "déc."))
    timeline = [(f"{month_names[int(month[5:]) - 1]} {month[:4]}", values[0], len(values[1]))
                for month, values in sorted(months.items())]
    table(graphs, 1, 31, [tr("Month", "Mois"), lines_label, requests_label], timeline)
    for col in range(25, 34):
        graphs.column_dimensions[get_column_letter(col)].hidden = True

    def text_style(size=11, bold=False):
        props = CharacterProperties(sz=size * 100, b=bold, solidFill=navy)
        return RichText(bodyPr=RichTextProperties(rot=0),
                        p=[Paragraph(pPr=ParagraphProperties(defRPr=props), endParaRPr=props)])

    def title_style(title):
        title.txPr = text_style(14, True)
        for paragraph in title.tx.rich.p:
            paragraph.pPr = ParagraphProperties(defRPr=CharacterProperties(sz=1400, b=True, solidFill=navy))
            for run in paragraph.r:
                run.rPr = CharacterProperties(sz=1400, b=True, solidFill=navy)

    def chart(cls, title, col, count, anchor, metric, values, value_offset=1):
        if not count:
            return
        obj = cls()
        if cls is BarChart:
            obj.x_axis.axPos = "l"
            obj.y_axis.axPos = "b"
            obj.type = "bar"
            obj.gapWidth = 65
        obj.title = title
        title_style(obj.title)
        obj.style = 13
        obj.width, obj.height = 18.5, 9.5
        obj.add_data(Reference(graphs, min_col=col + value_offset, min_row=1, max_row=count + 1), titles_from_data=True)
        categories = Reference(graphs, min_col=col, min_row=2, max_row=count + 1)
        # Explicit string categories keep localized month/material labels textual.
        obj.series[0].cat = AxDataSource(strRef=StrRef(f=str(categories)))
        obj.visible_cells_only = False
        obj.legend = None
        obj.graphical_properties = GraphicalProperties(solidFill="FFFFFF", ln=LineProperties(noFill=True))
        obj.plot_area.spPr = GraphicalProperties(solidFill="FFFFFF", ln=LineProperties(noFill=True))
        for axis in (obj.x_axis, obj.y_axis):
            axis.txPr = text_style()
            axis.tickLblPos = "nextTo"
            axis.majorTickMark = "none"
            axis.minorTickMark = "none"
            axis.spPr = GraphicalProperties(ln=LineProperties(solidFill="CAD5E0", w=9525))
            axis.majorGridlines = None
        obj.x_axis.tickLblSkip = 1
        obj.x_axis.tickMarkSkip = 1
        obj.y_axis.title = metric
        title_style(obj.y_axis.title)
        obj.y_axis.title.txPr = text_style(11)
        for paragraph in obj.y_axis.title.tx.rich.p:
            for run in paragraph.r:
                run.rPr = CharacterProperties(sz=1100, solidFill=navy)
        obj.y_axis.scaling.min = 0
        peak = max(values, default=0)
        # Modest headroom; small count series use integer ticks.
        if metric != quantity_label:
            step = max(1, ceil(peak / 5))
            obj.y_axis.majorUnit = step
            obj.y_axis.scaling.max = max(step, ceil((peak + 1) / step) * step)
            obj.y_axis.numFmt = "0"
        else:
            obj.y_axis.numFmt = "0.##"
        series = obj.series[0]
        color = "2874A6" if cls is BarChart else "178578"
        series.graphicalProperties.solidFill = color
        series.graphicalProperties.line.solidFill = color
        if cls is BarChart:
            obj.x_axis.scaling.orientation = "maxMin"
            obj.y_axis.crosses = "max"
            obj.x_axis.tickLblPos = "low"
            obj.layout = Layout(manualLayout=ManualLayout(layoutTarget="inner", xMode="factor", yMode="factor",
                                                        x=0.40, y=0.20, w=0.52, h=0.64))
            series.graphicalProperties.line.noFill = True
            series.graphicalProperties.line.solidFill = None
            obj.dataLabels = DataLabelList(showVal=True, showLegendKey=False, showSerName=False,
                                          showCatName=False, dLblPos="outEnd", txPr=text_style(), numFmt="0.##")
        else:
            obj.x_axis.axPos = "b"
            obj.y_axis.axPos = "l"
            obj.x_axis.title = tr("Month", "Mois")
            title_style(obj.x_axis.title)
            obj.x_axis.title.txPr = text_style(11)
            for paragraph in obj.x_axis.title.tx.rich.p:
                for run in paragraph.r:
                    run.rPr = CharacterProperties(sz=1100, solidFill=navy)
            obj.layout = Layout(manualLayout=ManualLayout(layoutTarget="inner", xMode="factor", yMode="factor",
                                                        x=0.15, y=0.20, w=0.78, h=0.62))
            obj.y_axis.majorGridlines = ChartLines(spPr=GraphicalProperties(ln=LineProperties(solidFill="E4EAF0", w=6350)))
            series.graphicalProperties.line.width = 25400
            series.marker.symbol = "circle"
            series.marker.size = 7
            series.marker.graphicalProperties.solidFill = color
            series.marker.graphicalProperties.line.solidFill = color
            series.smooth = False
            if count <= 6:
                obj.dataLabels = DataLabelList(showVal=True, showLegendKey=False, showSerName=False,
                                              showCatName=False, dLblPos="t", txPr=text_style(), numFmt="0")
        graphs.add_chart(obj, anchor)

    chart(BarChart, tr("Top 10 Materials Issued", "Top 10 des articles sortis"), 25, len(top), "A8", quantity_label, [m[3] for m in top])
    chart(BarChart, tr("Approved Material Requests by Department", "Demandes de matériel approuvées par département"),
          28, len(depts), "L8", requests_label, [len(d[1]) for d in depts])
    chart(LineChart, tr("Material Lines Issued Over Time", "Lignes d'articles sorties dans le temps"),
          31, len(timeline), "A27", lines_label, [t[1] for t in timeline])
    chart(LineChart, tr("Approved Material Requests Over Time", "Demandes de matériel approuvées au fil du temps"),
          31, len(timeline), "L27", requests_label, [t[2] for t in timeline], value_offset=2)
    graphs.print_area = "A1:U43"
    graphs.page_setup.fitToHeight = 1
    graphs.sheet_view.zoomScale = 85
