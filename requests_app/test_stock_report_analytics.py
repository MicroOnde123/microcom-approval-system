from datetime import date, datetime, timezone as dt_timezone
from io import BytesIO
from pathlib import Path
import os
from unittest import skipUnless
from unittest.mock import patch

from django.conf import settings
from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from accounts.models import Department
from inventory.models import Material, MaterialCategory
from .models import Request, RequestMaterialItem, RequestType


class StockReportAnalyticsTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.origin = Department.objects.create(name="Administration", code="AN-ADMIN")
        cls.destination = Department.objects.create(name="Technique", code="AN-TECH")
        cls.other = Department.objects.create(name="Other", code="AN-OTHER")
        cls.user = get_user_model().objects.create_user(
            username="analytics-stock", can_manage_stock=True, department=cls.origin)
        cls.request_type = RequestType.objects.create(name="Materials", code="AN-MAT", requires_materials=True)
        category = MaterialCategory.objects.create(name="Analytics", code="AN")
        cls.cable = Material.objects.create(name="Cable", code="AN-C", category=category, unit="m", stock_quantity=100)
        cls.bolt = Material.objects.create(name="Bolt", code="AN-B", category=category, unit="pcs", stock_quantity=100)
        cls.make_request("SEP-1", cls.destination, date(2026, 9, 1), [(cls.cable, 5), (cls.bolt, 2)])
        cls.make_request("SEP-2", cls.destination, date(2026, 9, 30), [(cls.cable, 3)])
        cls.make_request("SEP-OTHER", cls.other, date(2026, 9, 15), [(cls.bolt, 4)])
        cls.make_request("AUG", cls.destination, date(2026, 8, 31), [(cls.cable, 99)])
        cls.make_request("OCT", cls.destination, date(2026, 10, 1), [(cls.cable, 88)])
        cls.make_request("PENDING", cls.destination, date(2026, 9, 15), [(cls.cable, 77)], status="PENDING")

    @classmethod
    def make_request(cls, number, dept, needed, items, status="APPROVED"):
        req = Request.objects.create(request_number=number, request_type=cls.request_type,
            submitted_by=cls.user, department=cls.origin, request_for_department=dept,
            date_needed=needed, finalized_at=datetime(needed.year, needed.month, needed.day, tzinfo=dt_timezone.utc),
            status=status, description="Fixture", material_issue_note="Preserved note")
        for material, quantity in items:
            RequestMaterialItem.objects.create(request=req, material=material, quantity=quantity)
        return req

    def export(self, params=None, language="en"):
        self.client.force_login(self.user)
        response = self.client.get(reverse("export_material_report_excel"), params or {},
                                   HTTP_ACCEPT_LANGUAGE=language, HTTP_HOST="127.0.0.1")
        self.assertEqual(response.status_code, 200)
        return response, load_workbook(BytesIO(response.content))

    def test_first_sheet_matches_original_export_including_styles(self):
        params = {"date_from": "2026-09-01", "date_to": "2026-09-30", "department": self.destination.pk}
        with patch("requests_app.views.timezone.now", return_value=datetime(2026, 10, 5, tzinfo=dt_timezone.utc)):
            with patch("requests_app.stock_report_analytics.add_stock_report_analytics"):
                _, original = self.export(params)
            _, enhanced = self.export(params)
        self.assertEqual(enhanced.sheetnames, ["Material Report", "Dashboard", "Graphs"])
        before, after = original.active, enhanced.active
        self.assertEqual(list(before.values), list(after.values))
        for row in before:
            for cell in row:
                actual = after[cell.coordinate]
                self.assertEqual(cell._style, actual._style)
                self.assertEqual(cell.number_format, actual.number_format)
        self.assertEqual(str(before.merged_cells), str(after.merged_cells))
        self.assertEqual(before.freeze_panes, after.freeze_panes)
        self.assertEqual(before.auto_filter.ref, after.auto_filter.ref)
        self.assertEqual(before.page_setup, after.page_setup)
        self.assertEqual({k: dict(v) for k, v in before.column_dimensions.items()},
                         {k: dict(v) for k, v in after.column_dimensions.items()})
        rows = list(after.values)[4:]
        self.assertEqual(len(rows), 3)
        self.assertEqual({r[0] for r in rows}, {"SEP-1", "SEP-2"})
        self.assertEqual(sum(r[9] for r in rows), 10)

    def assert_analytics_match_rows(self, wb):
        rows = list(wb.active.values)[4:]
        dash, graphs = wb.worksheets[1:]
        self.assertEqual(dash["A10"].value, len({r[0] for r in rows}))
        self.assertEqual(dash["D9"].value, "Material Lines Issued")
        self.assertEqual(dash["D10"].value, len(rows))
        self.assertEqual(dash["G10"].value, len({r[7] for r in rows}))
        self.assertEqual(dash["J10"].value, len({r[3] for r in rows}))
        material, departments, department_lines, monthly, requests = {}, {}, {}, {}, {}
        for row in rows:
            label = f"{row[6]} ({row[10]})" if row[10] else row[6]
            material[label] = material.get(label, 0) + row[9]
            departments.setdefault(row[3], set()).add(row[0])
            department_lines[row[3]] = department_lines.get(row[3], 0) + 1
            month = row[5][:7]
            monthly[month] = monthly.get(month, 0) + 1
            requests.setdefault(month, set()).add(row[0])
        def pairs(col, size):
            return {graphs.cell(r, col).value: graphs.cell(r, col + 1).value for r in range(2, size + 2)}
        self.assertEqual(pairs(25, len(material)), material)
        self.assertEqual(pairs(28, len(departments)), {name: len(ids) for name, ids in departments.items()})
        friendly_months = {month: datetime.strptime(month, "%Y-%m").strftime("%b %Y") for month in monthly}
        self.assertEqual(pairs(31, len(monthly)), {friendly_months[m]: count for m, count in monthly.items()})
        self.assertEqual([graphs.cell(r, 31).value for r in range(2, len(monthly) + 2)],
                         [friendly_months[m] for m in sorted(monthly)])
        self.assertEqual([dash.cell(16, c).value for c in (6, 7, 8)], ["Department", "Approved Requests", "Material Lines"])
        self.assertEqual({dash.cell(r, 6).value: (dash.cell(r, 7).value, dash.cell(r, 8).value)
                          for r in range(17, 17 + len(departments))},
                         {name: (len(ids), department_lines[name]) for name, ids in departments.items()})
        self.assertEqual({dash.cell(r, 1).value: (dash.cell(r, 2).value, dash.cell(r, 3).value, dash.cell(r, 4).value)
                          for r in range(17, 17 + len(material))},
                         {row[6]: (row[7], row[10], sum(r[9] for r in rows if r[7] == row[7])) for row in rows})
        for r in range(2, len(monthly) + 2):
            month = sorted(monthly)[r - 2]
            self.assertEqual(graphs.cell(r, 33).value, len(requests[month]))
        self.assertEqual(len(graphs._charts), 4)
        self.assertTrue(all(not chart.visible_cells_only for chart in graphs._charts))

    def test_filters_and_aggregations_share_export_scope(self):
        for params in ({}, {"date_from": "2026-09-01", "date_to": "2026-09-30"},
                       {"department": self.other.pk}, {"q": "AN-C"}, {"q": "SEP-1"},
                       {"q": "analytics-stock"}, {"date_from": "2026-09-01"},
                       {"date_to": "2026-09-30"}):
            with self.subTest(params=params):
                _, wb = self.export(params)
                self.assert_analytics_match_rows(wb)
        # A material search selects requests, then exports ALL their items.
        _, wb = self.export({"q": "AN-C", "date_from": "2026-09-01", "date_to": "2026-09-30"})
        self.assertEqual(wb["Dashboard"]["D10"].value, 3)

    def test_empty_and_french_workbooks(self):
        for language, names, message in (("en", ["Material Report", "Dashboard", "Graphs"], "No stock data"),
                                         ("fr", ["Rapport de matériel", "Tableau de bord", "Graphiques"], "Aucune donnée")):
            _, wb = self.export({"q": "no-match"}, language)
            self.assertEqual(wb.sheetnames, names)
            self.assertEqual(wb.active.max_row, 4)
            self.assertEqual(wb.worksheets[1]["D10"].value, 0)
            for sheet in wb.worksheets[1:]:
                self.assertIn(message, sheet["A15"].value)
                self.assertEqual(len(sheet._charts), 0)
        _, wb = self.export(language="fr")
        self.assertEqual(wb.worksheets[1]["D9"].value, "Lignes d'articles sorties")
        self.assertEqual(wb.worksheets[2]["AE2"].value, "août 2026")
        self.assertEqual(wb.worksheets[2]["AE3"].value, "sept. 2026")
        self.assertEqual(wb.worksheets[2]["AE4"].value, "oct. 2026")
        expected_titles = ["Top 10 des articles sortis", "Demandes de matériel approuvées par département",
                           "Lignes d'articles sorties dans le temps", "Demandes de matériel approuvées au fil du temps"]
        self.assertEqual([c.title.tx.rich.p[0].r[0].t for c in wb.worksheets[2]._charts], expected_titles)
        self.assertEqual(wb.worksheets[2]._charts[1].y_axis.title.tx.rich.p[0].r[0].t, "Demandes approuvées")
        self.assertEqual(wb.worksheets[2]._charts[2].y_axis.title.tx.rich.p[0].r[0].t, "Lignes d'articles")
        self.assertEqual(wb.worksheets[2]._charts[2].x_axis.title.tx.rich.p[0].r[0].t, "Mois")

    def test_missing_approval_date_is_explicit_and_not_fabricated(self):
        Request.objects.filter(status="APPROVED").update(finalized_at=None)
        _, wb = self.export()
        self.assertEqual(len(wb["Graphs"]._charts), 2)
        self.assertIn("excluded from time charts: 5", wb["Dashboard"]["A12"].value)

    def test_chart_presentation_and_small_exports(self):
        for params, months in (({"q": "SEP-OTHER"}, 1),
                               ({"date_from": "2026-08-01", "date_to": "2026-09-30"}, 2)):
            with self.subTest(params=params):
                _, wb = self.export(params)
                self.assert_analytics_match_rows(wb)
                graph_sheet = wb["Graphs"]
                self.assertEqual(graph_sheet["A2"].value, "STOCK REPORT — VISUAL ANALYSIS")
                self.assertIn("same filtered records", graph_sheet["A5"].value)
                self.assertEqual([(c.anchor._from.col, c.anchor._from.row) for c in graph_sheet._charts],
                                 [(0, 7), (11, 7), (0, 26), (11, 26)])
                for chart in graph_sheet._charts:
                    self.assertEqual(chart.anchor.ext.cx, 6660000)  # 18.5 cm
                    self.assertEqual(chart.anchor.ext.cy, 3420000)  # 9.5 cm
                    self.assertEqual(chart.title.tx.rich.p[0].r[0].rPr.sz, 1400)
                    self.assertEqual(chart.x_axis.txPr.p[0].pPr.defRPr.sz, 1100)
                    self.assertEqual(chart.x_axis.txPr.bodyPr.rot, 0)
                    self.assertIsNone(chart.legend)
                    self.assertTrue(chart.dataLabels.showVal)
                    self.assertTrue(all(graph_sheet.column_dimensions[col].hidden for col in ("Y", "AB", "AE")))
                for chart in graph_sheet._charts[:2]:
                    self.assertEqual(chart.type, "bar")
                    self.assertEqual(chart.x_axis.scaling.orientation, "maxMin")
                    self.assertEqual(chart.dataLabels.dLblPos, "outEnd")
                for chart in graph_sheet._charts[2:]:
                    self.assertEqual(chart.series[0].marker.symbol, "circle")
                    self.assertEqual(chart.series[0].marker.size, 7)
                    self.assertEqual(chart.y_axis.scaling.min, 0)
                    self.assertGreaterEqual(chart.y_axis.majorUnit, 1)
                    self.assertEqual(chart.series[0].cat.strRef.f,
                                     f"'Graphs'!$AE$2:$AE${months + 1}" if months > 1 else "'Graphs'!$AE$2")
                self.assertEqual([c.title.tx.rich.p[0].r[0].t for c in graph_sheet._charts],
                                 ["Top 10 Materials Issued", "Approved Material Requests by Department",
                                  "Material Lines Issued Over Time", "Approved Material Requests Over Time"])

    def test_unitless_material_and_chronological_year_boundary(self):
        self.cable.unit = ""
        self.cable.save(update_fields=["unit"])
        self.make_request("DECEMBER", self.destination, date(2025, 12, 1), [(self.cable, 2)])
        self.make_request("JANUARY", self.destination, date(2026, 1, 1), [(self.cable, 3)])
        _, wb = self.export({"date_from": "2025-12-01", "date_to": "2026-01-31"})
        self.assert_analytics_match_rows(wb)
        self.assertEqual(wb["Graphs"]["Y2"].value, "Cable")
        self.assertEqual([wb["Graphs"].cell(r, 31).value for r in (2, 3)], ["Dec 2025", "Jan 2026"])
        _, french = self.export({"date_from": "2025-12-01", "date_to": "2026-01-31"}, "fr")
        self.assertEqual([french["Graphiques"].cell(r, 31).value for r in (2, 3)], ["déc. 2025", "janv. 2026"])

    def test_permissions(self):
        self.user.can_manage_stock = False
        self.user.save(update_fields=["can_manage_stock"])
        self.client.force_login(self.user)
        self.assertEqual(self.client.get(reverse("export_material_report_excel"), HTTP_HOST="127.0.0.1").status_code, 403)

    def test_queries_do_not_grow_per_request(self):
        from django.db import connection
        from django.test.utils import CaptureQueriesContext
        from django.test import RequestFactory
        from .views import export_material_report_excel
        request = RequestFactory().get("/", {"q": "SEP-1"})
        request.user = self.user
        with CaptureQueriesContext(connection) as small:
            export_material_report_excel(request)
        request = RequestFactory().get("/")
        request.user = self.user
        with CaptureQueriesContext(connection) as large:
            export_material_report_excel(request)
        self.assertEqual(len(small), len(large))

    @skipUnless(os.environ.get("GENERATE_STOCK_REPORT_SAMPLES") == "1", "Opt-in local sample generation")
    def test_generate_local_representative_workbooks(self):
        # Synthetic test fixtures only; no development/production data copied.
        third_department = Department.objects.create(name="Administration & Logistics", code="AN-LOG")
        extra_materials = [Material.objects.create(name=name, code=code, category=self.cable.category,
                                                  unit=unit, stock_quantity=500)
                           for name, code, unit in (("Printer Toner HP 59A", "AN-TON", "unit"),
                                                    ("RJ45 Connector Packs", "AN-RJ", "box"),
                                                    ("Network Router", "AN-RTR", "pcs"))]
        materials = [self.cable, self.bolt, *extra_materials]
        departments = [self.destination, self.other, third_department]
        for month in (8, 9, 10):
            for index in range(5):
                items = [(materials[(index + offset) % 5], (index + 1) * (month - 6) + offset)
                         for offset in range(1 + index % 3)]
                self.make_request(f"SAMPLE-{month}-{index}", departments[(month + index) % 3],
                                  date(2026, month, 3 + index * 5), items)
        folder = Path(settings.BASE_DIR) / "local_artifacts"
        folder.mkdir(exist_ok=True)
        for language in ("en", "fr"):
            response, wb = self.export({"date_from": "2026-08-01", "date_to": "2026-10-31"}, language)
            self.assertGreaterEqual(wb.worksheets[1]["G10"].value, 5)
            self.assertGreaterEqual(wb.worksheets[1]["J10"].value, 3)
            self.assertEqual(len(wb.worksheets[2]._charts), 4)
            (folder / f"stock_report_sample_{language}.xlsx").write_bytes(response.content)
