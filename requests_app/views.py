import csv
import time
import uuid
from decimal import Decimal, InvalidOperation
from django.http import HttpResponse
from urllib.parse import urlencode

from django.core.exceptions import ValidationError
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.urls import reverse
from django.utils.crypto import get_random_string
from django.contrib import messages
from django.db import models, transaction
from django.db.models.functions import RowNumber
from django.http import HttpResponseForbidden, JsonResponse
from django.utils import timezone
from django.views.decorators.http import require_POST

from .forms import MaterialIssueNoteForm, RequestForm, RequestMaterialItemFormSet
from .approval_forms import ApprovalActionForm
from .stock_forms import ReturnToStockForm
from .models import Request, RequestApproval, RequestAttachment, RequestMaterialItem, StockMovement, RequestAuditLog
from .services import submit_request, approve_step, reject_step, return_step, resubmit_request
from django.utils.translation import get_language, gettext as _

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from accounts.models import Department
from inventory.models import Material



MATERIAL_FORMSET_PREFIX = "material_items"
MATERIAL_TWO_COPY_ITEM_LIMIT = 6
MATERIAL_TWO_COPY_WARNING = (
    "This request has too many material items for two-copy printing. "
    "Please print one copy."
)
SUBMISSION_TOKEN_SESSION_KEY = "request_submission_tokens"
SUBMISSION_TOKEN_MAX_AGE = 24 * 60 * 60
SUBMISSION_TOKEN_LIMIT = 20


def _submission_tokens(request):
    now = int(time.time())
    tokens = request.session.get(SUBMISSION_TOKEN_SESSION_KEY, {})
    tokens = {
        token: data
        for token, data in tokens.items()
        if now - data.get("created_at", 0) <= SUBMISSION_TOKEN_MAX_AGE
    }
    return dict(
        sorted(
            tokens.items(),
            key=lambda item: item[1].get("created_at", 0),
            reverse=True,
        )[:SUBMISSION_TOKEN_LIMIT]
    )


def issue_submission_token(request, action):
    tokens = _submission_tokens(request)
    token = str(uuid.uuid4())
    tokens[token] = {
        "action": action,
        "status": "pending",
        "created_at": int(time.time()),
    }
    request.session[SUBMISSION_TOKEN_SESSION_KEY] = tokens
    request.session.modified = True
    return token


def submission_token_result(request, action):
    token = request.POST.get("submission_token", "")
    data = _submission_tokens(request).get(token)
    if not data or data.get("action") != action:
        return "invalid", None, token
    return data.get("status", "invalid"), data.get("request_id"), token


def mark_submission_token_used(request, token, request_id):
    tokens = _submission_tokens(request)
    if token in tokens:
        tokens[token]["status"] = "used"
        tokens[token]["request_id"] = request_id
        request.session[SUBMISSION_TOKEN_SESSION_KEY] = tokens
        request.session.modified = True


def safe_next_url(raw_url, default_url):
    if raw_url and raw_url.startswith("/") and not raw_url.startswith("//"):
        return raw_url

    return default_url


def current_path_with_query(request):
    querystring = request.GET.urlencode()

    if querystring:
        return f"{request.path}?{querystring}"

    return request.path


def requires_materials(req):
    return req.request_type and req.request_type.requires_materials


def request_type_behavior_context(form):
    request_type_field = form.fields["request_type"]
    queryset = request_type_field.queryset

    return {
        str(request_type.id): {
            "is_permission_request": request_type.is_permission_request,
            "requires_materials": request_type.requires_materials,
            "requires_amount": request_type.requires_amount,
        }
        for request_type in queryset
    }


def save_attachments(request, req):
    files = request.FILES.getlist("attachments")

    for f in files:
        RequestAttachment.objects.create(
            request=req,
            file=f,
            original_name=f.name,
            uploaded_by=request.user,
        )

def build_permission_metadata(form):
    cleaned = form.cleaned_data

    return {
        "permission_group": cleaned.get("permission_group"),
        "permission_subgroup": cleaned.get("permission_subgroup"),
        "destination": cleaned.get("destination"),
        "exit_reason": cleaned.get("exit_reason"),
        "departure_time": (
            cleaned.get("departure_time").strftime("%H:%M")
            if cleaned.get("departure_time")
            else None
        ),
        "return_time": (
            cleaned.get("return_time").strftime("%H:%M")
            if cleaned.get("return_time")
            else None
        ),
        "arrival_time": (
            cleaned.get("arrival_time").strftime("%H:%M")
            if cleaned.get("arrival_time")
            else None
        ),
        "driver_name": cleaned.get("driver_name"),
        "site": cleaned.get("site"),
        "valid_from": (
            cleaned.get("valid_from").strftime("%Y-%m-%d")
            if cleaned.get("valid_from")
            else None
        ),
        "valid_to": (
            cleaned.get("valid_to").strftime("%Y-%m-%d")
            if cleaned.get("valid_to")
            else None
        ),
        "microcom_agents": cleaned.get("microcom_agents"),
        "tt": cleaned.get("tt"),
        "external_persons": cleaned.get("external_persons"),
    }


def request_form_context(form, formset, submission_token, **extra):
    context = {
        "form": form,
        "formset": formset,
        "request_type_behavior": request_type_behavior_context(form),
        "submission_token": submission_token,
    }
    context.update(extra)
    return context


def submitted_material_count(formset):
    count = 0
    material_ids = []

    for form in formset.forms:
        if not hasattr(form, "cleaned_data"):
            continue

        cleaned = form.cleaned_data
        if not cleaned or cleaned.get("DELETE"):
            continue

        material = cleaned.get("material")
        quantity = cleaned.get("quantity")
        if material and quantity:
            count += 1
            material_ids.append(material.id)

    return count, len(material_ids) != len(set(material_ids))


def save_complete_draft_material_rows(req, post_data):
    prefix = MATERIAL_FORMSET_PREFIX
    total = int(post_data.get(f"{prefix}-TOTAL_FORMS") or 0)
    initial = int(post_data.get(f"{prefix}-INITIAL_FORMS") or 0)
    seen_material_ids = set()

    for index in range(total):
        material_id = post_data.get(f"{prefix}-{index}-material")
        quantity_value = post_data.get(f"{prefix}-{index}-quantity")
        note = post_data.get(f"{prefix}-{index}-note", "")
        delete_requested = post_data.get(f"{prefix}-{index}-DELETE")
        item_id = post_data.get(f"{prefix}-{index}-id")
        existing_item = None

        if item_id:
            existing_item = req.material_items.filter(id=item_id).first()

        if delete_requested and existing_item:
            existing_item.delete()
            continue

        if not material_id or not quantity_value:
            continue

        try:
            material = Material.objects.get(id=material_id, is_active=True)
            quantity = Decimal(quantity_value)
        except (Material.DoesNotExist, InvalidOperation):
            continue

        if quantity <= 0 or material.id in seen_material_ids:
            continue

        seen_material_ids.add(material.id)

        if existing_item and index < initial:
            existing_item.material = material
            existing_item.quantity = quantity
            existing_item.note = note
            existing_item.save()
        else:
            RequestMaterialItem.objects.create(
                request=req,
                material=material,
                quantity=quantity,
                note=note,
            )


def save_valid_material_formset(formset):
    formset.save()

@login_required
def create_request(request):
    if not request.user.department:
        messages.error(request, "Your account is not linked to any department.")
        return redirect("dashboard")

    if request.method == "POST":
        action = request.POST.get("action", "submit")
        is_draft_action = action == "save_draft"
        token_status, existing_request_id, submission_token = submission_token_result(
            request, "create"
        )
        if token_status == "used" and existing_request_id:
            messages.info(request, _("This request was already saved."))
            return redirect("request_detail", request_id=existing_request_id)
        if token_status != "pending":
            messages.error(request, _("This form has expired. Please submit it again."))
            return redirect("create_request")

        form = RequestForm(
            request.POST,
            request.FILES,
            user=request.user,
            draft=is_draft_action,
        )
        formset = RequestMaterialItemFormSet(
            request.POST,
            prefix=MATERIAL_FORMSET_PREFIX,
        )

        if form.is_valid():
            req = form.save(commit=False)
            req.request_number = f"REQ-{get_random_string(6).upper()}"
            req.submitted_by = request.user
            req.department = request.user.department
            req.metadata_json = build_permission_metadata(form)
            req.status = "DRAFT" if is_draft_action else "PENDING"

            material_request = requires_materials(req)

            if is_draft_action:
                req.save()
                if material_request:
                    save_complete_draft_material_rows(req, request.POST)
                save_attachments(request, req)
                mark_submission_token_used(request, submission_token, req.id)
                messages.success(request, _("Draft saved successfully."))
                return redirect("request_detail", request_id=req.id)

            if material_request:
                if not formset.is_valid():
                    return render(
                        request,
                        "requests_app/create_request.html",
                        request_form_context(form, formset, submission_token),
                    )

                material_count, has_duplicates = submitted_material_count(formset)

                if has_duplicates:
                    messages.error(request, _("Duplicate material rows are not allowed."))
                    return render(
                        request,
                        "requests_app/create_request.html",
                        request_form_context(form, formset, submission_token),
                    )

                if material_count <= 0:
                    messages.error(request, _("At least one material is required for material requests."))
                    return render(
                        request,
                        "requests_app/create_request.html",
                        request_form_context(form, formset, submission_token),
                    )

            req.save()

            if material_request:
                formset.instance = req
                save_valid_material_formset(formset)

            save_attachments(request, req)

            try:
                submit_request(req)
            except Exception:
                req.delete()
                messages.error(request, _("No approval workflow is configured for this request type."))
                return render(
                    request,
                    "requests_app/create_request.html",
                    request_form_context(form, formset, submission_token),
                )

            mark_submission_token_used(request, submission_token, req.id)
            messages.success(request, _("Request submitted successfully."))
            return redirect("dashboard")

    else:
        form = RequestForm(user=request.user)
        formset = RequestMaterialItemFormSet(prefix=MATERIAL_FORMSET_PREFIX)
        submission_token = issue_submission_token(request, "create")

    return render(
        request,
        "requests_app/create_request.html",
        request_form_context(form, formset, submission_token),
    )

def is_stock_manager(user):
    return (
        user.is_superuser
        or 
            getattr(user, "can_manage_stock", False)
    )
    
@login_required
def my_requests(request):
    qs = Request.objects.filter(submitted_by=request.user).order_by("-submitted_at")
    department = request.GET.get("department", "").strip()
    if department:
        qs = qs.filter(request_for_department_id=department)
    return render(request, "requests_app/my_requests.html", {
        "requests": qs.select_related("request_for_department"),
        "departments": Department.objects.order_by("name"),
        "selected_department": department,
    })


@login_required
def pending_approvals(request):
    approvals = RequestApproval.objects.filter(
        models.Q(approver_user=request.user)
        | models.Q(alternate_approver_user=request.user),
        status="PENDING",
        request__current_step_order=models.F("step_order"),
        request__status__in=["PENDING", "IN_REVIEW"],
    ).select_related(
        "request",
        "request__request_type",
        "request__request_for_department",
    ).order_by("-request__submitted_at", "-created_at")

    date_from = request.GET.get("date_from", "").strip()
    date_to = request.GET.get("date_to", "").strip()
    request_type = request.GET.get("request_type", "").strip()
    q = request.GET.get("q", "").strip()
    department = request.GET.get("department", "").strip()

    if date_from:
        approvals = approvals.filter(request__date_needed__gte=date_from)

    if date_to:
        approvals = approvals.filter(request__date_needed__lte=date_to)

    if request_type:
        approvals = approvals.filter(request__request_type_id=request_type)

    if department:
        approvals = approvals.filter(request__request_for_department_id=department)

    if q:
        approvals = approvals.filter(
            models.Q(request__request_number__icontains=q)
            | models.Q(request__submitted_by__username__icontains=q)
            | models.Q(request__submitted_by__full_name__icontains=q)
            | models.Q(request__submitted_by__email__icontains=q)
    )

    request_types = Request.objects.values_list(
        "request_type__id",
        "request_type__name",
    ).distinct()
    today = timezone.now().date()
    
    return render(
        request,
        "requests_app/pending_approvals.html",
        {
            "approvals": approvals,
            "request_types": request_types,
            "date_from": date_from,
            "date_to": date_to,
            "selected_request_type": request_type,
            "today": today,
            "q": q,
            "departments": Department.objects.order_by("name"),
            "selected_department": department,
            "current_list_url": current_path_with_query(request),
        },
    )


@login_required
def approval_detail(request, approval_id):
    back_url = safe_next_url(
        request.POST.get("next") or request.GET.get("next"),
        reverse("pending_approvals"),
    )
    approval = get_object_or_404(
        RequestApproval.objects.select_related(
            "request",
            "request__request_type",
            "approver_user",
            "alternate_approver_user",
        ).filter(
            models.Q(approver_user=request.user)
            | models.Q(alternate_approver_user=request.user)
        ),
        id=approval_id,
    )

    if request.method != "POST" and (
        approval.status != "PENDING"
        or approval.request.current_step_order != approval.step_order
        or approval.request.status not in ["PENDING", "IN_REVIEW"]
    ):
        messages.error(request, "This approval is no longer active.")
        return redirect(back_url)

    if request.method == "POST":
        form = ApprovalActionForm(request.POST)

        if form.is_valid():
            action = form.cleaned_data["action"]
            comment = form.cleaned_data["comment"]

            try:
                if action == "APPROVE":
                    approve_step(approval, request.user, comment)
                    messages.success(request, "Request approved successfully.")
                elif action == "REJECT":
                    reject_step(approval, request.user, comment)
                    messages.success(request, "Request rejected successfully.")
                elif action == "RETURN":
                    return_step(approval, request.user, comment)
                    messages.success(request, "Request returned for changes.")
            except Exception as e:
                messages.error(request, str(e))
                return redirect(
                    f"{reverse('approval_detail', args=[approval.id])}?{urlencode({'next': back_url})}"
                )

            return redirect(back_url)
    else:
        form = ApprovalActionForm()

    return render(
        request,
        "requests_app/approval_detail.html",
        {"approval": approval, "request_obj": approval.request, "form": form, "back_url": back_url},
    )


@login_required
def request_detail(request, request_id):
    request_obj = get_object_or_404(
        Request.objects.prefetch_related(
            "approvals__approver_user",
            "approvals__acted_by",
            "attachments",
            "audit_logs",
            "material_items__material__category",
        ),
        id=request_id,
    )

    is_submitter = request_obj.submitted_by == request.user
    is_approver = request_obj.approvals.filter(
        models.Q(approver_user=request.user)
        | models.Q(alternate_approver_user=request.user)
    ).exists()
    is_stock_user = is_stock_manager(request.user) and request_obj.material_items.exists()
    
    if not is_submitter and not is_approver and not request.user.is_superuser and not is_stock_user:
        return HttpResponseForbidden("You are not allowed to view this request.")

    if is_submitter:
        default_back_url = reverse("my_requests")
    elif is_stock_user:
        default_back_url = reverse("material_reports")
    elif is_approver:
        default_back_url = reverse("pending_approvals")
    else:
        default_back_url = reverse("dashboard")

    back_url = safe_next_url(request.GET.get("next"), default_back_url)

    active_approval_for_user = None
    if request_obj.status in ["PENDING", "IN_REVIEW"]:
        active_approval_for_user = request_obj.approvals.filter(
            models.Q(approver_user=request.user)
            | models.Q(alternate_approver_user=request.user),
            status="PENDING",
            step_order=request_obj.current_step_order,
        ).first()

    can_cancel_request = (
        request_obj.status in ["DRAFT", "RETURNED", "PENDING", "IN_REVIEW"]
        and (is_submitter or request.user.is_superuser)
    )
    can_edit_request = is_submitter and request_obj.status in ["DRAFT", "RETURNED"]

    return render(
        request,
        "requests_app/request_detail.html",
        {
            "request_obj": request_obj,
            "back_url": back_url,
            "active_approval_for_user": active_approval_for_user,
            "can_cancel_request": can_cancel_request,
            "can_edit_request": can_edit_request,
            "can_edit_material_issue_note": (
                request_obj.status == "APPROVED"
                and requires_materials(request_obj)
                and is_stock_manager(request.user)
            ),
            "material_issue_note_form": MaterialIssueNoteForm(instance=request_obj),
        },
    )


@login_required
@require_POST
def update_material_issue_note(request, request_id):
    request_obj = get_object_or_404(
        Request.objects.select_related("request_type"),
        id=request_id,
    )
    detail_url = reverse("request_detail", args=[request_obj.id])
    back_url = safe_next_url(request.POST.get("next"), detail_url)

    if not is_stock_manager(request.user):
        return HttpResponseForbidden(
            _("You are not allowed to edit the material issue note.")
        )

    if request_obj.status != "APPROVED":
        messages.error(
            request,
            _("Only approved requests can have a material issue note."),
        )
        return redirect(back_url)

    if not requires_materials(request_obj):
        messages.error(
            request,
            _("Only material requests can have a material issue note."),
        )
        return redirect(back_url)

    form = MaterialIssueNoteForm(request.POST, instance=request_obj)
    if not form.is_valid():
        messages.error(request, _("The material issue note could not be saved."))
        return redirect(back_url)

    with transaction.atomic():
        form.save()
        RequestAuditLog.objects.create(
            request=request_obj,
            action="MATERIAL_ISSUE_NOTE_UPDATED",
            performed_by=request.user,
            comment=_("Material issue note updated."),
        )

    messages.success(request, _("Material issue note saved successfully."))
    return redirect(back_url)


@login_required
def edit_request(request, request_id):
    request_obj = get_object_or_404(Request, id=request_id, submitted_by=request.user)

    if request_obj.status not in ["DRAFT", "RETURNED"]:
        messages.error(request, _("Only draft or returned requests can be edited."))
        return redirect("request_detail", request_id=request_obj.id)

    submission_action = f"edit:{request_obj.id}"
    if request.method == "POST":
        action = request.POST.get("action", "resubmit")
        if action not in {"save_draft", "submit", "save_changes", "resubmit"}:
            action = "resubmit"
        is_draft_save = action in {"save_draft", "save_changes"}
        is_submit_action = action in {"submit", "resubmit"}
        token_status, existing_request_id, submission_token = submission_token_result(
            request, submission_action
        )
        if token_status == "used" and existing_request_id:
            messages.info(request, _("This request was already saved."))
            return redirect("request_detail", request_id=existing_request_id)
        if token_status != "pending":
            messages.error(request, _("This form has expired. Please submit it again."))
            return redirect("edit_request", request_id=request_obj.id)

        form = RequestForm(
            request.POST,
            request.FILES,
            instance=request_obj,
            user=request.user,
            draft=is_draft_save,
        )

        if form.is_valid():
            req = form.save(commit=False)
            req.department = request.user.department
            req.metadata_json = build_permission_metadata(form)

            material_request = requires_materials(req)

            formset = RequestMaterialItemFormSet(
                request.POST,
                instance=req,
                prefix=MATERIAL_FORMSET_PREFIX,
            )

            if is_draft_save:
                req.save()
                if material_request:
                    save_complete_draft_material_rows(req, request.POST)
                else:
                    req.material_items.all().delete()

                save_attachments(request, req)
                mark_submission_token_used(request, submission_token, req.id)
                if request_obj.status == "DRAFT":
                    messages.success(request, _("Draft saved successfully."))
                else:
                    messages.success(request, _("Changes saved successfully."))
                return redirect("request_detail", request_id=req.id)

            if is_submit_action and material_request:
                if not formset.is_valid():
                    return render(
                        request,
                        "requests_app/edit_request.html",
                        request_form_context(
                            form,
                            formset,
                            submission_token,
                            request_obj=request_obj,
                            back_url=reverse("request_detail", args=[request_obj.id]),
                        ),
                    )

                material_count, has_duplicates = submitted_material_count(formset)

                if has_duplicates:
                    messages.error(request, _("Duplicate material rows are not allowed."))
                    return render(
                        request,
                        "requests_app/edit_request.html",
                        request_form_context(
                            form,
                            formset,
                            submission_token,
                            request_obj=request_obj,
                            back_url=reverse("request_detail", args=[request_obj.id]),
                        ),
                    )

                if material_count <= 0:
                    messages.error(request, _("At least one material is required for material requests."))
                    return render(
                        request,
                        "requests_app/edit_request.html",
                        request_form_context(
                            form,
                            formset,
                            submission_token,
                            request_obj=request_obj,
                            back_url=reverse("request_detail", args=[request_obj.id]),
                        ),
                    )
            else:
                formset = RequestMaterialItemFormSet(
                    instance=req,
                    prefix=MATERIAL_FORMSET_PREFIX,
                )

            req.save()

            if material_request:
                save_valid_material_formset(formset)
            else:
                req.material_items.all().delete()

            save_attachments(request, req)

            try:
                if request_obj.status == "DRAFT":
                    req.status = "PENDING"
                    req.current_step_order = None
                    req.finalized_at = None
                    req.save()
                    submit_request(req)
                    success_message = _("Request submitted successfully.")
                else:
                    resubmit_request(req, request.user)
                    success_message = _("Request updated and resubmitted successfully.")
            except ValidationError as exc:
                messages.error(request, exc.messages[0] if hasattr(exc, "messages") else str(exc))
                return render(
                    request,
                    "requests_app/edit_request.html",
                    request_form_context(
                        form,
                        formset,
                        submission_token,
                        request_obj=request_obj,
                        back_url=reverse("request_detail", args=[request_obj.id]),
                    ),
                )

            mark_submission_token_used(request, submission_token, req.id)
            messages.success(request, success_message)
            return redirect("request_detail", request_id=req.id)

        formset = RequestMaterialItemFormSet(
            request.POST,
            instance=request_obj,
            prefix=MATERIAL_FORMSET_PREFIX,
        )

    else:
        form = RequestForm(instance=request_obj, user=request.user)
        formset = RequestMaterialItemFormSet(
            instance=request_obj,
            prefix=MATERIAL_FORMSET_PREFIX,
        )
        submission_token = issue_submission_token(request, submission_action)

    return render(
        request,
        "requests_app/edit_request.html",
        request_form_context(
            form,
            formset,
            submission_token,
            request_obj=request_obj,
            back_url=reverse("request_detail", args=[request_obj.id]),
        ),
    )


@login_required
@require_POST
def cancel_request(request, request_id):
    request_obj = get_object_or_404(Request, id=request_id)
    detail_url = reverse("request_detail", args=[request_obj.id])
    back_url = safe_next_url(request.POST.get("next"), detail_url)

    if not (request_obj.submitted_by == request.user or request.user.is_superuser):
        return HttpResponseForbidden(_("You are not allowed to cancel this request."))

    if request_obj.status not in ["DRAFT", "RETURNED", "PENDING", "IN_REVIEW"]:
        messages.error(request, _("This request cannot be cancelled."))
        return redirect(back_url)

    with transaction.atomic():
        request_obj = Request.objects.select_for_update().get(id=request_obj.id)

        if request_obj.status not in ["DRAFT", "RETURNED", "PENDING", "IN_REVIEW"]:
            messages.error(request, _("This request cannot be cancelled."))
            return redirect(back_url)

        request_obj.status = "CANCELLED"
        request_obj.current_step_order = None
        request_obj.finalized_at = timezone.now()
        request_obj.save(update_fields=["status", "current_step_order", "finalized_at"])

        RequestAuditLog.objects.create(
            request=request_obj,
            action="CANCELLED",
            performed_by=request.user,
            comment=_("Request cancelled by user."),
        )

    messages.success(request, _("Request cancelled successfully."))
    return redirect(back_url)


@login_required
def approved_document(request, request_id):
    copies = request.GET.get("copies", "1")
    if copies not in {"1", "2"}:
        copies = "1"

    request_obj = get_object_or_404(
        Request.objects.prefetch_related(
            "material_items__material",
            "approvals__approver_user",
            "approvals__acted_by",
        ),
        id=request_id,
    )

    allowed = (
        request.user == request_obj.submitted_by
        or request_obj.approvals.filter(
            models.Q(approver_user=request.user)
            | models.Q(alternate_approver_user=request.user)
        ).exists()
        or request.user.is_superuser
        or (
            is_stock_manager(request.user)
            and request_obj.material_items.exists()
        )
    )

    if not allowed:
        return HttpResponseForbidden("You are not allowed to view this document.")

    if request_obj.status != "APPROVED":
        messages.error(request, "This request is not approved yet.")
        return redirect("request_detail", request_id=request_obj.id)

    return render_approved_document(request, request_obj)


def render_approved_document(request, request_obj, document_url=None, default_back_url=None):
    copies = request.GET.get("copies", "1")
    if copies not in {"1", "2"}:
        copies = "1"
    two_copy_warning = ""
    if copies == "2" and request_obj.material_items.count() > MATERIAL_TWO_COPY_ITEM_LIMIT:
        copies = "1"
        two_copy_warning = MATERIAL_TWO_COPY_WARNING

    approvals = request_obj.approvals.filter(
        status="APPROVED"
    ).order_by("step_order")

    return render(
        request,
        "requests_app/approved_document.html",
        {
            "request_obj": request_obj,
            "approvals": approvals,
            "back_url": safe_next_url(
                request.GET.get("next"),
                default_back_url or reverse("request_detail", args=[request_obj.id]),
            ),
            "document_url": document_url or reverse("approved_document", args=[request_obj.id]),
            "copies": copies,
            "two_copy_warning": two_copy_warning,
        },
    )

@login_required
def permission_document(request, request_id):
    request_obj = get_object_or_404(
        Request.objects.prefetch_related(
            "approvals__approver_user",
            "approvals__acted_by",
        ),
        id=request_id,
    )

    allowed = (
        request.user == request_obj.submitted_by
        or request_obj.approvals.filter(
            models.Q(approver_user=request.user)
            | models.Q(alternate_approver_user=request.user)
        ).exists()
        or request.user.is_superuser
    )

    if not allowed:
        return HttpResponseForbidden("You are not allowed to view this document.")

    if request_obj.status != "APPROVED":
        messages.error(request, "This request is not approved yet.")
        return redirect("request_detail", request_id=request_obj.id)

    metadata = request_obj.metadata_json or {}

    if not metadata.get("permission_group"):
        messages.error(request, "This request is not a permission request.")
        return redirect("request_detail", request_id=request_obj.id)

    approvals = request_obj.approvals.filter(
        status="APPROVED"
    ).order_by("step_order")

    return render(
        request,
        "requests_app/permission_document.html",
        {
            "request_obj": request_obj,
            "metadata": metadata,
            "approvals": approvals,
            "back_url": safe_next_url(
                request.GET.get("next"),
                reverse("request_detail", args=[request_obj.id]),
            ),
        },
    )

@login_required
def approval_history(request):
    approvals = RequestApproval.objects.filter(
        models.Q(acted_by=request.user)
        | models.Q(acted_by__isnull=True, approver_user=request.user),
    ).exclude(
        status="PENDING"
    ).exclude(
        request__status__in=["DRAFT", "CANCELLED"]
    ).annotate(
        request_action_rank=models.Window(
            expression=RowNumber(),
            partition_by=[models.F("request_id")],
            order_by=[
                models.F("acted_at").desc(nulls_last=True),
                models.F("id").desc(),
            ],
        )
    ).filter(
        request_action_rank=1
    ).select_related(
        "request",
        "request__request_type",
        "request__request_for_department",
        "acted_by",
        "approver_user",
        "alternate_approver_user",
    ).order_by("-acted_at")

    department = request.GET.get("department", "").strip()
    if department:
        approvals = approvals.filter(request__request_for_department_id=department)

    return render(
        request,
        "requests_app/approval_history.html",
        {
            "approvals": approvals,
            "departments": Department.objects.order_by("name"),
            "selected_department": department,
            "current_list_url": current_path_with_query(request),
        },
    )
    
def get_filtered_material_report_requests(request):
    requests = Request.objects.filter(
        status="APPROVED",
        material_items__isnull=False,
    ).distinct().prefetch_related(
        "material_items__material__category",
        "approvals__approver_user",
        "approvals__acted_by",
    ).select_related(
        "submitted_by",
        "department",
        "request_for_department",
        "request_type",
    )

    q = request.GET.get("q", "").strip()
    date_from = request.GET.get("date_from", "").strip()
    date_to = request.GET.get("date_to", "").strip()
    department = request.GET.get("department", "").strip()

    if q:
        requests = requests.filter(
            models.Q(request_number__icontains=q)
            | models.Q(submitted_by__username__icontains=q)
            | models.Q(submitted_by__full_name__icontains=q)
            | models.Q(material_items__material__name__icontains=q)
            | models.Q(material_items__material__code__icontains=q)
        )

    if date_from:
        requests = requests.filter(date_needed__gte=date_from)

    if date_to:
        requests = requests.filter(date_needed__lte=date_to)

    if department:
        requests = requests.filter(request_for_department_id=department)

    return requests.order_by("-finalized_at", "-submitted_at")

@login_required
def material_reports(request):
    if not is_stock_manager(request.user):
        return HttpResponseForbidden("You are not allowed to access material reports.")

    requests = Request.objects.filter(
        status="APPROVED",
        material_items__isnull=False,
       
    ).distinct().prefetch_related(
        "material_items__material__category",
        "approvals__approver_user",
        "approvals__acted_by",
    ).select_related(
        "submitted_by",
        "department",
        "request_for_department",
        "request_type",
    )

    q = request.GET.get("q", "").strip()
    date_from = request.GET.get("date_from", "").strip()
    date_to = request.GET.get("date_to", "").strip()
    department = request.GET.get("department", "").strip()

    if q:
        requests = requests.filter(
            models.Q(request_number__icontains=q)
            | models.Q(submitted_by__username__icontains=q)
            | models.Q(submitted_by__full_name__icontains=q)
            | models.Q(material_items__material__name__icontains=q)
            | models.Q(material_items__material__code__icontains=q)
        )

    if date_from:
        requests = requests.filter(date_needed__gte=date_from)

    if date_to:
        requests = requests.filter(date_needed__lte=date_to)

    if department:
        requests = requests.filter(request_for_department_id=department)

    departments = Department.objects.order_by("name").values_list("id", "name")

    requests = requests.order_by("-finalized_at", "-submitted_at")

    return render(
        request,
        "requests_app/material_reports.html",
        {
            "requests": requests,
            "departments": departments,
            "q": q,
            "date_from": date_from,
            "date_to": date_to,
            "selected_department": department,
            "current_list_url": current_path_with_query(request),
            "active_querystring": request.GET.urlencode(),
        },
    )

@login_required
def export_material_report_csv(request):
    if not is_stock_manager(request.user):
        return HttpResponseForbidden("You are not allowed to export material reports.")

    requests = Request.objects.filter(
        status="APPROVED",
        material_items__isnull=False,
    ).distinct().prefetch_related(
        "material_items__material__category",
        "approvals__approver_user",
        "approvals__acted_by",
    ).select_related(
        "submitted_by",
        "department",
        "request_for_department",
        "request_type",
    )

    q = request.GET.get("q", "").strip()
    date_from = request.GET.get("date_from", "").strip()
    date_to = request.GET.get("date_to", "").strip()
    department = request.GET.get("department", "").strip()

    if q:
        requests = requests.filter(
            models.Q(request_number__icontains=q)
            | models.Q(submitted_by__username__icontains=q)
            | models.Q(submitted_by__full_name__icontains=q)
            | models.Q(material_items__material__name__icontains=q)
            | models.Q(material_items__material__code__icontains=q)
        )

    if date_from:
        requests = requests.filter(date_needed__gte=date_from)

    if date_to:
        requests = requests.filter(date_needed__lte=date_to)

    if department:
        requests = requests.filter(request_for_department_id=department)

    response = HttpResponse(content_type="text/csv")
    response["Content-Disposition"] = 'attachment; filename="material_report.csv"'

    writer = csv.writer(response)

    writer.writerow([
        _("Request Number"),
        _("Requester"),
        _("Department"),
        _("Request For Department"),
        _("Date Needed"),
        _("Approved Date"),
        _("Material"),
        _("Material Code"),
        _("Category"),
        _("Quantity"),
        _("Unit"),
        _("Available Stock"),
        _("Description"),
        _("Material Issue Note"),
        _("Approvers"),
    ])

    for req in requests.order_by("-finalized_at", "-submitted_at"):
        approvers = ", ".join([
            approval.display_approver_name
            for approval in req.approvals.all()
            if approval.status == "APPROVED"
        ])

        for item in req.material_items.all():
            writer.writerow([
                req.request_number,
                req.submitted_by.full_name or req.submitted_by.username,
                req.department.name if req.department else "",
                req.request_for_department.name,
                req.date_needed,
                req.finalized_at.strftime("%Y-%m-%d %H:%M") if req.finalized_at else "",
                item.material.name,
                item.material.code,
                item.material.category.name if item.material.category else "",
                item.quantity,
                item.material.unit,
                item.material.stock_quantity,
                req.description,
                req.material_issue_note,
                approvers,
            ])

    return response
    
@login_required
def bulk_print_material_documents(request):
    if not is_stock_manager(request.user):
        return HttpResponseForbidden("You are not allowed to bulk print material documents.")

    copies = request.GET.get("copies", "1")
    if copies not in {"1", "2"}:
        copies = "1"

    if request.method != "POST":
        return redirect("material_reports")

    back_url = safe_next_url(request.POST.get("next"), reverse("material_reports"))

    selected_ids = request.POST.getlist("selected_requests")

    if not selected_ids:
        messages.error(request, "Select at least one material request to print.")
        return redirect(back_url)

    requests = Request.objects.filter(
        id__in=selected_ids,
        status="APPROVED",
        material_items__isnull=False,
        
    ).distinct().prefetch_related(
        "material_items__material",
        "approvals__approver_user",
        "approvals__acted_by",
    ).select_related(
        "submitted_by",
        "department",
        "request_for_department",
    ).order_by("-finalized_at", "-submitted_at")

    two_copy_warning = ""
    if copies == "2":
        for req in requests:
            req.print_copies = "2"
            if req.material_items.count() > MATERIAL_TWO_COPY_ITEM_LIMIT:
                req.print_copies = "1"
                two_copy_warning = MATERIAL_TWO_COPY_WARNING
    else:
        for req in requests:
            req.print_copies = "1"

    return render(
        request,
        "requests_app/bulk_print_material_documents.html",
        {
            "requests": requests,
            "back_url": back_url,
            "copies": copies,
            "two_copy_warning": two_copy_warning,
        },
    )

@login_required
def return_material_to_stock(request, item_id):
    if not is_stock_manager(request.user):
        return HttpResponseForbidden(_("You are not allowed to return stock."))

    item = get_object_or_404(
        RequestMaterialItem.objects.select_related(
            "request",
            "material",
        ),
        id=item_id,
        request__status="APPROVED",
    )

    req = item.request
    material = item.material
    back_url = safe_next_url(
        request.POST.get("next") or request.GET.get("next"),
        reverse("material_reports"),
    )

    if request.method == "POST":
        form = ReturnToStockForm(request.POST)

        if form.is_valid():
            quantity = form.cleaned_data["quantity"]
            reason = form.cleaned_data["reason"]

            if quantity > item.quantity:
                messages.error(
                    request,
                    _("Returned quantity cannot be greater than requested quantity."),
                )
                return redirect("return_material_to_stock", item_id=item.id)

            material.stock_quantity += quantity
            material.save(update_fields=["stock_quantity"])

            StockMovement.objects.create(
                material=material,
                request=req,
                quantity=quantity,
                movement_type="RETURN",
                performed_by=request.user,
                note=_("Material returned to stock."),
                return_reason=reason,
            )

            RequestAuditLog.objects.create(
                request=req,
                action="STOCK_RETURNED",
                performed_by=request.user,
                comment=_("Returned %(quantity)s %(unit)s of %(material)s to stock. Reason: %(reason)s")
                % {
                    "quantity": quantity,
                    "unit": material.unit or "",
                    "material": material.name,
                    "reason": reason,
                },
            )

            messages.success(
                request,
                _("Material returned to stock successfully."),
            )

            return redirect(back_url)

    else:
        form = ReturnToStockForm()

    return render(
        request,
        "requests_app/return_material_to_stock.html",
        {
            "form": form,
            "item": item,
            "request_obj": req,
            "material": material,
            "back_url": back_url,
        },
    )

@login_required
def notification_count(request):
    pending_count = RequestApproval.objects.filter(
        models.Q(approver_user=request.user)
        | models.Q(alternate_approver_user=request.user),
        status="PENDING",
        request__current_step_order=models.F("step_order"),
        request__status__in=["PENDING", "IN_REVIEW"],
    ).count()

    returned_count = Request.objects.filter(
        submitted_by=request.user,
        status="RETURNED",
    ).count()

    return JsonResponse({
        "pending": pending_count,
        "returned": returned_count,
    })


@login_required
def export_material_report_excel(request):
    if not is_stock_manager(request.user):
        return HttpResponseForbidden(
            _("You are not allowed to export material reports.")
        )

    requests = get_filtered_material_report_requests(request)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = _("Material Report")

    headers = [
        _("Request Number"),
        _("Requester"),
        _("Department"),
        _("Request For Department"),
        _("Date Needed"),
        _("Approved Date"),
        _("Material"),
        _("Material Code"),
        _("Category"),
        _("Quantity"),
        _("Unit"),
        _("Available Stock"),
        _("Description"),
        _("Material Issue Note"),
        _("Approvers"),
    ]

    sheet.merge_cells("A1:O1")
    sheet["A1"] = _("Microcom Material Report")
    sheet["A1"].font = Font(bold=True, size=14)
    sheet["A1"].alignment = Alignment(horizontal="center")

    sheet["A2"] = _("Generated At")
    sheet["B2"] = timezone.now().strftime("%Y-%m-%d %H:%M")

    start_row = 4

    for col, header in enumerate(headers, start=1):
        cell = sheet.cell(row=start_row, column=col, value=header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    row_num = start_row + 1

    for req in requests:
        approvers = ", ".join([
            approval.display_approver_name
            for approval in req.approvals.all()
            if approval.status == "APPROVED"
        ])

        for item in req.material_items.all():
            sheet.append([
                req.request_number,
                req.submitted_by.full_name or req.submitted_by.username,
                req.department.name if req.department else "",
                req.request_for_department.name,
                req.date_needed,
                req.finalized_at.strftime("%Y-%m-%d %H:%M") if req.finalized_at else "",
                item.material.name,
                item.material.code,
                item.material.category.name if item.material.category else "",
                item.quantity,
                item.material.unit,
                item.material.stock_quantity,
                req.description,
                req.material_issue_note,
                approvers,
            ])
            row_num += 1

    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    for row in sheet.iter_rows(min_row=start_row, max_row=sheet.max_row, min_col=1, max_col=len(headers)):
        for cell in row:
            cell.border = thin_border
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    sheet.freeze_panes = "A5"
    sheet.auto_filter.ref = f"A{start_row}:O{sheet.max_row}"

    widths = {
        "A": 18,
        "B": 24,
        "C": 20,
        "D": 24,
        "E": 15,
        "F": 20,
        "G": 30,
        "H": 18,
        "I": 22,
        "J": 12,
        "K": 12,
        "L": 16,
        "M": 45,
        "N": 45,
        "O": 35,
    }

    for col, width in widths.items():
        sheet.column_dimensions[col].width = width

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    filename = (
        "rapport_materiel.xlsx"
        if get_language() == "fr"
        else "material_report.xlsx"
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'

    workbook.save(response)
    return response


def administration_report_scope():
    from django.conf import settings
    return Request.objects.filter(
        status="APPROVED",
        request_for_department__code=settings.ADMINISTRATION_REPORT_DEPARTMENT_CODE,
        request_type__code=settings.ADMINISTRATION_REPORT_REQUEST_TYPE_CODE,
        request_type__requires_materials=False,
        request_type__is_permission_request=False,
        material_items__isnull=True,
    ).select_related("submitted_by", "department", "request_for_department", "request_type")


def can_view_administration_reports(user):
    return user.is_authenticated and (user.is_superuser or user.can_view_administration_reports)


@login_required
def administration_reports(request):
    from django.core.paginator import Paginator
    from django.utils.dateparse import parse_date
    if not can_view_administration_reports(request.user):
        return HttpResponseForbidden(_("You are not allowed to access administration reports."))
    scope = administration_report_scope()
    rows = scope
    filters = {key: request.GET.get(key, "").strip() for key in
               ("q", "department", "requester", "date_from", "date_to")}
    if filters["q"]:
        query = models.Q()
        for field in ("request_number", "description", "submitted_by__username",
                      "submitted_by__full_name", "submitted_by__first_name", "submitted_by__last_name"):
            query |= models.Q(**{field + "__icontains": filters["q"]})
        rows = rows.filter(query)
    for key, field in (("department", "department_id"), ("requester", "submitted_by_id")):
        if filters[key]:
            if filters[key].isascii() and filters[key].isdigit() and len(filters[key]) < 19:
                rows = rows.filter(**{field: int(filters[key])})
            else:
                rows = rows.none()
    for key, lookup in (("date_from", "submitted_at__date__gte"), ("date_to", "submitted_at__date__lte")):
        if filters[key]:
            try:
                value = parse_date(filters[key])
            except ValueError:
                value = None
            rows = rows.filter(**{lookup: value}) if value else rows.none()
    params = request.GET.copy()
    params.pop("page", None)
    page = Paginator(rows.order_by("-finalized_at", "-submitted_at", "-pk"), 25).get_page(request.GET.get("page"))
    return render(request, "requests_app/administration_reports.html", {
        **filters, "requests": page, "page_obj": page,
        "departments": scope.order_by("department__name").values_list("department_id", "department__name").distinct(),
        "requesters": scope.order_by("submitted_by__username").values_list("submitted_by_id", "submitted_by__full_name", "submitted_by__username").distinct(),
        "current_list_url": current_path_with_query(request),
        "active_querystring": params.urlencode(),
    })


@login_required
def administration_report_document(request, request_id):
    if not can_view_administration_reports(request.user):
        return HttpResponseForbidden(_("You are not allowed to access administration reports."))
    request_obj = get_object_or_404(administration_report_scope().prefetch_related(
        "material_items__material", "approvals__approver_user", "approvals__acted_by"), pk=request_id)
    return render_approved_document(
        request, request_obj,
        document_url=reverse("administration_report_document", args=[request_obj.pk]),
        default_back_url=reverse("administration_reports"),
    )
